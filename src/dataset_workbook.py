"""Text-only dataset workbooks shared by the CLI and offline preparation tools."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
import math
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

INPUT_HEADERS = ("Natural Language", "Requirement ID")
EVALUATION_HEADERS = (*INPUT_HEADERS, "Raw Formula", "Normalized Formula")
COUNTS = {"A72": 72, "B240": 240, "C1026": 1026}


def normalized_requirement_id(value: object) -> str:
    if not isinstance(value, str) or not re.fullmatch(
        r"(?:[A-Za-z][A-Za-z0-9_-]*:)?[A-Za-z][A-Za-z0-9_-]*-\d+", value
    ):
        raise ValueError(f"无效的 Requirement ID: {value!r}")
    return value.replace(":", "-")


def _text(cell, *, location: str) -> str:
    if cell.data_type == "f":
        raise ValueError(f"{location}: 数据单元格应为文本，不能是 Excel 公式")
    if not isinstance(cell.value, str) or not cell.value.strip():
        raise ValueError(f"{location}: 需要非空文本")
    return cell.value


def read_source_sheet(sheet) -> list[tuple[str, str]]:
    """Read only NL and ID; reference columns are never selected here."""
    rows = sheet.iter_rows(min_col=1, max_col=2)
    header = next(rows, ())
    if tuple(cell.value for cell in header) != INPUT_HEADERS:
        raise ValueError(f"{sheet.title}: 前两列必须为 {INPUT_HEADERS}")
    result = []
    seen = set()
    for index, cells in enumerate(rows, start=2):
        if all(cell.value is None for cell in cells):
            continue
        location = f"{sheet.title} 第 {index} 行"
        text, source_id = (_text(cell, location=location) for cell in cells)
        normalized = normalized_requirement_id(source_id)
        if normalized in seen:
            raise ValueError(f"{location}: 重复的 Requirement ID: {source_id}")
        seen.add(normalized)
        result.append((source_id, text))
    if not result:
        raise ValueError(f"{sheet.title}: 没有有效需求")
    return result


def read_input_workbook(path: str | Path) -> list[tuple[str, str]]:
    """Accept one source-only sheet, rejecting evaluation workbooks as inputs."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except BadZipFile as error:
        raise ValueError("无法读取 XLSX：文件损坏或不是有效的 Excel 工作簿") from error
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("在线输入 XLSX 必须只有一个需求工作表；请使用 data/input 下的输入表")
        sheet = workbook.worksheets[0]
        if sheet.max_column != 2:
            raise ValueError("在线输入 XLSX 只允许 Natural Language 和 Requirement ID 两列，不接受参考答案列")
        return read_source_sheet(sheet)
    finally:
        workbook.close()


def workbook_bytes(sheets: dict[str, list[tuple]], *, evaluation: bool = False) -> bytes:
    """Build styled XLSX bytes with fixed document and ZIP timestamps."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    fixed_time = datetime(2000, 1, 1)
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time
    workbook.properties.creator = "Traj-LTLf"
    workbook.properties.title = "LTLf Evaluation Dataset" if evaluation else "LTLf Requirements"
    headers = EVALUATION_HEADERS if evaluation else INPUT_HEADERS
    widths = (100, 26, 76, 76) if evaluation else (100, 26)
    for number, (name, rows) in enumerate(sheets.items(), start=1):
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "C2"
        sheet.append(headers)
        for row in rows:
            if len(row) != len(headers) or any(not isinstance(value, str) or not value.strip() for value in row):
                raise ValueError(f"invalid workbook row in {name}")
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                # Keep NL and LTLf strings literal, including strings beginning '='.
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(name="Calibri", size=11, color="243746")
            line_count = max(
                sum(max(1, math.ceil(len(line) / (width * 0.85))) for line in value.split('\n'))
                for value, width in zip(row, widths)
            )
            sheet.row_dimensions[sheet.max_row].height = 16 * line_count + 12
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="175D66")
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 30
        table = Table(displayName=f"Dataset{number}", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    temporary = BytesIO()
    workbook.save(temporary)
    workbook.close()
    # openpyxl updates modified on save; normalize it and archive entry times.
    result = BytesIO()
    with ZipFile(temporary) as source, ZipFile(result, "w", ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)", rb"\g<1>2000-01-01T00:00:00Z\2", content)
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, content)
    return result.getvalue()
