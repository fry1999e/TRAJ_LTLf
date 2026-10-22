"""Offline deterministic XLSX input exports; no model or API calls."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from dataset_workbook import COUNTS, read_source_sheet, workbook_bytes

DATA_ROOT = ROOT / "data"
EVALUATION_FILE = DATA_ROOT / "LTLf_Evaluation_Dataset.xlsx"


def sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def build_exports(evaluation_file: Path = EVALUATION_FILE) -> dict[str, bytes]:
    """Select only NL/ID columns from the offline source; omit all reference columns."""
    exports = {}
    manifest = {"format": "manual-xlsx-input-v1", "datasets": {}}
    workbook = load_workbook(evaluation_file, read_only=True, data_only=False, keep_links=False)
    try:
        if workbook.sheetnames != list(COUNTS):
            raise ValueError("evaluation workbook must contain A72, B240, C1026 in that order")
        for dataset, count in COUNTS.items():
            source_rows = read_source_sheet(workbook[dataset])
            if len(source_rows) != count or any(not source_id.startswith(dataset + ":") for source_id, _ in source_rows):
                raise ValueError(f"wrong dataset IDs or count: {dataset}")
            content = workbook_bytes({"Requirements": [(text, source_id) for source_id, text in source_rows]})
            exports[f"{dataset}.xlsx"] = content
            manifest["datasets"][dataset] = {
                "count": count,
                "source": "data/LTLf_Evaluation_Dataset.xlsx",
                "sheet": dataset,
                "selected_columns": ["Natural Language", "Requirement ID"],
                "input": f"data/input/{dataset}.xlsx",
                "input_sha256": sha256(content),
                "ap_profile": "undetermined",
                "export_policy": "source_text_only; AP definitions require human review",
                "records": [{
                    "requirement_id": source_id.replace(":", "-"),
                    "qualified_id": source_id,
                    "requirement_sha256": sha256(text.encode("utf-8")),
                } for source_id, text in source_rows],
            }
    finally:
        workbook.close()
    exports["manifest.json"] = (json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    return exports


def verify_exports() -> None:
    for name, content in build_exports().items():
        target = DATA_ROOT / "input" / name
        if not target.is_file() or target.read_bytes() != content:
            raise ValueError(f"input export missing or modified: {target.relative_to(ROOT)}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--write", action="store_true", help="create missing exports; refuse changed files")
    group.add_argument("--verify", action="store_true", help="rebuild in memory and compare exact bytes")
    args = parser.parse_args()
    if args.write:
        exports = build_exports()
        for name, content in exports.items():
            target = DATA_ROOT / "input" / name
            if target.exists() and target.read_bytes() != content:
                raise ValueError(f"refusing to overwrite modified input: {name}")
        (DATA_ROOT / "input").mkdir(exist_ok=True)
        for name, content in exports.items():
            (DATA_ROOT / "input" / name).write_bytes(content)
    else:
        verify_exports()
    print("PASS: deterministic XLSX inputs A72=72, B240=240, C1026=1026 (1338 total)")


if __name__ == "__main__":
    main()
