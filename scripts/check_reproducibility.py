"""Offline dataset, input-loader and dependency verification; no API calls."""
from __future__ import annotations

import argparse
import hashlib
import importlib
from importlib.metadata import version, PackageNotFoundError
import json
from pathlib import Path
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from prepare_inputs import COUNTS, DATA_ROOT, EVALUATION_FILE, verify_exports
from dataset_workbook import EVALUATION_HEADERS, read_source_sheet
from requirement_loader import load_requirement_records
from runtime_config import load_runtime_config


def verify_data() -> dict:
    manifest = json.loads((DATA_ROOT / "manifest.json").read_text(encoding="utf-8"))
    if manifest.get("format") != "traj-ltlf-xlsx-data-v1":
        raise ValueError("unsupported data manifest")
    expected_files = {f"input/{dataset}.xlsx" for dataset in COUNTS}
    expected_files.update({EVALUATION_FILE.name, "input/example.xlsx"})
    if set(manifest["files"]) != expected_files:
        raise ValueError("missing or unexpected dataset files")
    for relative, expected in manifest["files"].items():
        path = (DATA_ROOT / relative).resolve()
        if not path.is_relative_to(DATA_ROOT.resolve()):
            raise ValueError(f"invalid dataset path: {relative}")
        if hashlib.sha256(path.read_bytes()).hexdigest() != expected:
            raise ValueError(f"dataset digest mismatch: {relative}")
    raw_root = (DATA_ROOT / "input" / "sources").resolve()
    for relative, expected in manifest["raw_files"].items():
        path = (raw_root / relative).resolve()
        if not path.is_relative_to(raw_root):
            raise ValueError(f"invalid source path: {relative}")
        if hashlib.sha256(path.read_bytes()).hexdigest() != expected:
            raise ValueError(f"source digest mismatch: {relative}")
    verify_exports()
    evaluation_records = []
    workbook = load_workbook(EVALUATION_FILE, read_only=True, data_only=False, keep_links=False)
    try:
        for dataset, count in COUNTS.items():
            records = load_requirement_records(DATA_ROOT / "input" / f"{dataset}.xlsx")
            sheet = workbook[dataset]
            if tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))) != EVALUATION_HEADERS:
                raise ValueError(f"invalid evaluation header: {dataset}")
            sources = read_source_sheet(sheet)
            if len(records) != count or [(r.requirement_id, r.text) for r in records] != [(source_id.replace(":", "-"), text) for source_id, text in sources]:
                raise ValueError(f"loader changed text, count or IDs: {dataset}")
            if manifest["datasets"][dataset] != {"count": count, "input": f"input/{dataset}.xlsx", "sheet": dataset}:
                raise ValueError(f"invalid dataset metadata: {dataset}")
            for cells in sheet.iter_rows(min_row=2):
                if len(cells) != 4 or any(cell.data_type == "f" or not isinstance(cell.value, str) or not cell.value.strip() for cell in cells):
                    raise ValueError(f"invalid evaluation row: {dataset}")
                text, source_id, raw, normalized = [cell.value for cell in cells]
                evaluation_records.append({"qualified_id": source_id, "requirement": text, "raw_formula": raw, "normalized_formula": normalized})
    finally:
        workbook.close()
    canonical = json.dumps(evaluation_records, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    if hashlib.sha256(canonical).hexdigest() != manifest["evaluation_records_sha256"]:
        raise ValueError("evaluation content changed after migration")
    return {"dataset_files": len(expected_files), "raw_files": len(manifest["raw_files"]), "requirements": sum(COUNTS.values())}


def check_dependencies() -> list[dict]:
    results = []
    for line in (ROOT / "requirements.txt").read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        name, expected = line.split("==")
        try:
            actual = version(name)
        except PackageNotFoundError:
            actual = "missing"
        results.append({"package": name, "expected": expected, "actual": actual, "match": actual == expected})
    return results


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check-dependencies", action="store_true", help="also require the pinned direct dependencies")
    args = parser.parse_args()
    if sys.version_info < (3, 13):
        raise SystemExit("Python 3.13+ is required")
    config = load_runtime_config()
    for name in ("cli", "pipeline", "batch_pipeline", "pre_ir_review", "contracts.dataset_semantics_contract"):
        importlib.import_module(name)
    data = verify_data()
    report = {"status": "pass", "python": sys.version.split()[0], "mode": config["experiment_mode"], **data}
    if args.check_dependencies:
        report["dependencies"] = check_dependencies()
        if not all(item["match"] for item in report["dependencies"]):
            report["status"] = "fail"
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if report["status"] != "pass":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
